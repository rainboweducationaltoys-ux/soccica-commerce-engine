import os
import re
import tempfile
import zipfile
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

from database import get_connection, initialize_database, utc_now
from logging_setup import configure_logging
from product_store import migrate_product_table


logger = configure_logging("soccica.spotlight_importer")


BASE_DIR = Path(__file__).resolve().parent.parent

SPOTLIGHT_EXPORTS_DIR = (
    BASE_DIR
    / "input"
    / "spotlight_exports"
)

REQUIRED_COLUMNS = [
    "ASIN",
    "Domain",
    "Link",
    "Title",
    "Category",
    "Brand",
    "Sellers List",
]

# ASIN'ler bu durumlardaysa spotlight ile yeniden fiyatlandırmak güvenli —
# henüz Shopify'a gönderilmemiş veya zaten spotlight akışında.
SAFE_TO_REPRICE_STATUSES = {
    None,
    "pending",
    "retry",
    "error",
    "spotlight_priced",
    "spotlight_retry",
}

AMAZON_FBA_OFFER_PATTERN = re.compile(
    r"^FBA\s*\|\s*Amazon\s+AED\s*([\d.]+)\s*\(\+AED\s*([\d.]+)\)"
)


def safe_text(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def find_newest_export_file() -> Path:
    files = sorted(
        SPOTLIGHT_EXPORTS_DIR.glob("*.xlsx"),
        key=lambda file_path: file_path.stat().st_mtime,
        reverse=True,
    )

    if not files:
        raise FileNotFoundError(
            "input/spotlight_exports/ içinde .xlsx dosyası "
            f"bulunamadı: {SPOTLIGHT_EXPORTS_DIR}"
        )

    return files[0]


def _patch_sheet_state(path: Path) -> Path:
    file_descriptor, patched_path_text = tempfile.mkstemp(
        suffix=".xlsx"
    )
    os.close(file_descriptor)

    patched_path = Path(patched_path_text)

    with zipfile.ZipFile(path, "r") as source_zip:
        with zipfile.ZipFile(
            patched_path,
            "w",
        ) as target_zip:
            for item in source_zip.infolist():
                data = source_zip.read(
                    item.filename
                )

                if item.filename == "xl/workbook.xml":
                    data = data.replace(
                        b'state="show"',
                        b'state="visible"',
                    )

                target_zip.writestr(item, data)

    return patched_path


def _exception_chain(
    error: BaseException,
) -> list[BaseException]:
    chain: list[BaseException] = []
    current: BaseException | None = error

    while current is not None:
        chain.append(current)
        current = current.__cause__

    return chain


def load_spotlight_workbook(path: Path):
    path = Path(path)

    try:
        return openpyxl.load_workbook(
            path,
            data_only=True,
            read_only=True,
        )
    except ValueError as error:
        combined_text = " ".join(
            str(current)
            for current in _exception_chain(error)
        ).lower()

        if (
            "state" not in combined_text
            and "visible" not in combined_text
            and "hidden" not in combined_text
        ):
            raise

        logger.warning(
            "%s açılırken sheet state hatası alındı, "
            "xl/workbook.xml yamanıp yeniden deneniyor: %s",
            path.name,
            error,
        )

        patched_path = _patch_sheet_state(path)

        return openpyxl.load_workbook(
            patched_path,
            data_only=True,
            read_only=True,
        )


def read_spotlight_rows(
    path: Path,
) -> list[dict[str, Any]]:
    workbook = load_spotlight_workbook(path)
    sheet = workbook.active

    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter, None)

    if not header_row:
        return []

    header_index = {
        safe_text(name): index
        for index, name in enumerate(header_row)
        if safe_text(name)
    }

    missing_columns = [
        column
        for column in REQUIRED_COLUMNS
        if column not in header_index
    ]

    if missing_columns:
        raise ValueError(
            "Spotlight export dosyasında beklenen sütunlar "
            "eksik: " + ", ".join(missing_columns)
        )

    def cell(
        raw_row: tuple[Any, ...],
        column: str,
    ) -> Any:
        index = header_index[column]

        if index >= len(raw_row):
            return None

        return raw_row[index]

    rows: list[dict[str, Any]] = []

    for raw_row in rows_iter:
        if raw_row is None:
            continue

        rows.append(
            {
                "asin": safe_text(
                    cell(raw_row, "ASIN")
                ).upper(),
                "domain": safe_text(
                    cell(raw_row, "Domain")
                ).lower(),
                "link": safe_text(
                    cell(raw_row, "Link")
                ),
                "title": safe_text(
                    cell(raw_row, "Title")
                ),
                "category": safe_text(
                    cell(raw_row, "Category")
                ),
                "brand": safe_text(
                    cell(raw_row, "Brand")
                ),
                "sellers_list": safe_text(
                    cell(raw_row, "Sellers List")
                ),
            }
        )

    return rows


def parse_amazon_fba_offers(
    sellers_list_text: str,
) -> list[tuple[float, float]]:
    offers: list[tuple[float, float]] = []

    for line in sellers_list_text.splitlines():
        match = AMAZON_FBA_OFFER_PATTERN.match(
            line.strip()
        )

        if not match:
            continue

        price = float(match.group(1))
        shipping = float(match.group(2))

        offers.append((price, shipping))

    return offers


def select_best_offer_index(
    offers: list[tuple[float, float]],
) -> int | None:
    if not offers:
        return None

    return min(
        range(len(offers)),
        key=lambda index: offers[index][0],
    )


def save_spotlight_offers(
    asin: str,
    offers: list[tuple[float, float]],
    selected_index: int,
    scan_date: str,
    source_file: str,
) -> None:
    with get_connection() as connection:
        for index, (price, shipping) in enumerate(
            offers
        ):
            connection.execute(
                """
                INSERT INTO spotlight_offers (
                    asin,
                    amazon_price_aed,
                    amazon_shipping_aed,
                    is_selected,
                    scan_date,
                    source_file
                )
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    asin,
                    price,
                    shipping,
                    1 if index == selected_index else 0,
                    scan_date,
                    source_file,
                ),
            )


def upsert_spotlight_product(
    asin: str,
    title: str,
    brand: str,
    category: str,
    link: str,
    amazon_price_aed: float,
    scan_date: str,
) -> str:
    now = utc_now()

    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT status
            FROM products
            WHERE asin = ?
            """,
            (asin,),
        ).fetchone()

        if row is None:
            connection.execute(
                """
                INSERT INTO products (
                    asin,
                    status,
                    title,
                    brand,
                    source_url,
                    spotlight_amazon_price_aed,
                    spotlight_category,
                    spotlight_scan_date,
                    created_at,
                    updated_at
                )
                VALUES (
                    ?, 'spotlight_priced', ?, ?, ?, ?, ?, ?, ?, ?
                )
                """,
                (
                    asin,
                    title,
                    brand,
                    link,
                    amazon_price_aed,
                    category,
                    scan_date,
                    now,
                    now,
                ),
            )

            return "inserted"

        current_status = row["status"]

        if current_status in SAFE_TO_REPRICE_STATUSES:
            connection.execute(
                """
                UPDATE products
                SET
                    status = 'spotlight_priced',
                    title = ?,
                    brand = ?,
                    source_url = ?,
                    spotlight_amazon_price_aed = ?,
                    spotlight_category = ?,
                    spotlight_scan_date = ?,
                    updated_at = ?
                WHERE asin = ?
                """,
                (
                    title,
                    brand,
                    link,
                    amazon_price_aed,
                    category,
                    scan_date,
                    now,
                    asin,
                ),
            )

            return "repriced"

        connection.execute(
            """
            UPDATE products
            SET
                spotlight_amazon_price_aed = ?,
                spotlight_category = ?,
                spotlight_scan_date = ?,
                updated_at = ?
            WHERE asin = ?
            """,
            (
                amazon_price_aed,
                category,
                scan_date,
                now,
                asin,
            ),
        )

        return "updated_only"


def import_spotlight_export(
    path: Path | None = None,
) -> dict[str, int]:
    if path is None:
        path = find_newest_export_file()

    scan_date = date.today().isoformat()
    source_file = path.name

    print(
        "Spotlight export dosyası okunuyor:",
        path,
    )

    rows = read_spotlight_rows(path)

    stats = {
        "rows_read": 0,
        "non_ae_skipped": 0,
        "ineligible": 0,
        "inserted": 0,
        "repriced": 0,
        "updated_only": 0,
    }

    for row in rows:
        stats["rows_read"] += 1

        asin = row["asin"]

        if not asin:
            continue

        if row["domain"] and row["domain"] != "ae":
            stats["non_ae_skipped"] += 1

            logger.warning(
                "%s: beklenmeyen Domain değeri (%s), "
                "satır atlandı.",
                asin,
                row["domain"],
            )

            continue

        offers = parse_amazon_fba_offers(
            row["sellers_list"]
        )

        selected_index = select_best_offer_index(
            offers
        )

        if selected_index is None:
            stats["ineligible"] += 1
            continue

        save_spotlight_offers(
            asin=asin,
            offers=offers,
            selected_index=selected_index,
            scan_date=scan_date,
            source_file=source_file,
        )

        selected_price, _selected_shipping = offers[
            selected_index
        ]

        action = upsert_spotlight_product(
            asin=asin,
            title=row["title"],
            brand=row["brand"],
            category=row["category"],
            link=row["link"],
            amazon_price_aed=selected_price,
            scan_date=scan_date,
        )

        stats[action] += 1

    logger.info(
        "Spotlight import tamamlandı (%s): %s",
        source_file,
        stats,
    )

    return stats


def main() -> None:
    initialize_database()
    migrate_product_table()

    stats = import_spotlight_export()

    print("=" * 70)
    print("SPOTLIGHT IMPORT TAMAMLANDI")
    print("=" * 70)

    for key, value in stats.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
